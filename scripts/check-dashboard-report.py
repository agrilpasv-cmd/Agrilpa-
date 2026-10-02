"""Read-only independent checks of reports from test-dashboard-report.cjs."""
import json, sys, zipfile, posixpath
from pathlib import Path
from datetime import datetime, timedelta
from xml.etree import ElementTree as ET
import openpyxl

root = Path(sys.argv[1])
ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
manifest = json.loads((root/'manifest.json').read_text(encoding='utf-8'))
for case in manifest:
    file = root/case['file']
    data, role = case['data'], case['role']
    view = data['views'][role]
    with zipfile.ZipFile(file) as z:
        assert z.testzip() is None, file
        names = set(z.namelist())
        for name in names:
            tree = ET.fromstring(z.read(name))
            if name.endswith('.rels'):
                folder = posixpath.dirname(posixpath.dirname(name))
                for relationship in tree:
                    assert relationship.get('TargetMode') != 'External'
                    target = posixpath.normpath(posixpath.join(folder,relationship.get('Target'))).lstrip('/')
                    assert target in names,(name,target)
            if '/charts/chart' in name:
                for reference in tree.findall('.//c:numRef',ns):
                    assert reference.find('c:f',ns) is not None
                    assert reference.find('c:numCache/c:ptCount',ns) is not None
    values = openpyxl.load_workbook(file,data_only=True)
    formulas = openpyxl.load_workbook(file,data_only=False)
    assert values.sheetnames == ['Resumen','Gráficos','Actividad diaria','Totales','Pedidos','Clasificación']
    summary = values['Resumen']
    metrics = {r[0]:r for r in summary.iter_rows(min_row=11,values_only=True) if r[0] and r[1]}
    for metric in view['metrics']+view['secondary']:
        if metric['id']=='amount': continue
        row = metrics[metric['label']]
        assert row[2] == ('Sin datos' if metric['value'] is None else metric['value'])
        if metric['previous'] is not None:
            assert row[3] == metric['previous']
            if metric['value'] is not None:
                assert abs(row[4]-(metric['value']-metric['previous'])) < 1e-9
    currencies = set(view['money']['current']) | set(view['money']['previous'])
    for currency in currencies:
        row = next(r for r in summary.iter_rows(values_only=True) if r[0]=='Importe de pedidos completados' and r[1]==currency)
        assert row[2]==view['money']['current'].get(currency,0)
        assert row[3]==view['money']['previous'].get(currency,0)
    daily = list(values['Actividad diaria'].iter_rows(min_row=6,values_only=True))
    assert len(daily)==sum(len(s['points']) for s in view['series'])
    for series in view['series']:
        rows=[r for r in daily if r[0]==series['label']]
        for row, point in zip(rows,series['points']):
            assert row[1].date().isoformat()==point['date']
            assert row[3].date()==row[1].date()-timedelta(days=data['period']['days'])
            assert row[2]==point['current'] and row[4]==point['previous']
        totals=next(r for r in values['Totales'].iter_rows(min_row=6,values_only=True) if r[0]==series['label'])
        assert totals[1]==sum(p['current'] for p in series['points'])
        assert totals[2]==sum(p['previous'] for p in series['points'])
    assert values['Pedidos'].cell(6+len(view['orders']),2).value==sum(o['value'] for o in view['orders'])
    charts=formulas['Gráficos']._charts
    assert len(charts)==len(view['series'])+1+bool(view['ranking'])
    for chart in charts:
        for series in chart.series:
            reference=series.val.numRef
            sheet_name, area=reference.f.rsplit('!',1)
            sheet_name=sheet_name.strip("'")
            worksheet_values=[cell.value for row in values[sheet_name][area] for cell in row]
            cached_values=[p.v for p in reference.numCache.pt]
            assert worksheet_values==cached_values,(case['file'],reference.f)
    for sheet in formulas:
        assert not sheet.sheet_view.showGridLines
        for row in sheet:
            for cell in row:
                assert cell.data_type!='e',(file,cell.coordinate,cell.value)
                if cell.data_type=='f': assert not cell.value.startswith('=HYPERLINK')
    if case['file']=='unavailable.xlsx':
        assert metrics['Conversaciones nuevas'][2]=='Sin datos'
        assert metrics['Respondidas en 24 horas'][2]=='Sin datos'
    assert summary['A1'].font.name=='Segoe UI'
print(f'Validated {len(manifest)} XLSX files: CRC, XML relationships, typed dates/numbers, comparisons, multi-currency totals, formulas and editable chart data.')
